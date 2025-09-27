import os
from datetime import datetime

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv('.secure')


class PositionExtractor:
    """Class PositionExtractor contains method for extracting position from
    Google"""

    def __init__(self, file_name, url_service, headers):
        self.file_name = file_name
        self.url_service = url_service
        self.headers = headers

    def load_data_from_excel(self):
        """Method loads data from excel file"""
        wb = load_workbook(self.file_name)
        ws = wb.active
        return wb, ws

    @staticmethod
    def create_payload(device, language_code, target_url, keyword):
        """Method creates payload"""
        payload = [{
            "keyword": keyword,
            "location_code": 2804,
            "language_code": language_code,
            "device": device,
            "os": "windows" if device == "desktop" else "android",
            "depth": 50,
            "target": target_url
        }]
        return payload

    def send_request(self, payload):
        """Method sends Post request and returns response wish position
        information"""
        response = requests.post(self.url_service, headers=self.headers,
                                 json=payload)
        return response

    @staticmethod
    def parse_response(response):
        """Method parses response from web service"""
        if response.status_code != 200:
            print(response.status_code)
            return "Not Found"
        data = response.json()
        task = data.get("tasks", [])[0]
        if task["status_code"] != 20000:
            print(f"Not_found_status_code: {task['status_code']}")
            return "Not in TOP"
        result = task["result"][0]
        if result.get("items") and isinstance(result["items"], list) and len(
                result["items"]) > 0:
            rank = result["items"][0].get("rank_group", None)
            print(f"position {rank}")
            return rank if rank is not None else "Not in TOP"
        else:
            return "Not Found"

    def get_site_position_in_file(self):
        wb, ws = self.load_data_from_excel()
        next_row = ws.max_row + 1
        df = pd.read_excel(self.file_name, header=None)
        ws.cell(row=next_row, column=1,
                value=datetime.now().strftime('%Y/%m/%d %H-%M-%S'))

        for col in range(2, df.shape[1] + 1):
            try:
                device = str(df.iloc[0, col - 1]).strip().lower()
                language_code = str(df.iloc[1, col - 1]).strip().lower()
                target_url = str(df.iloc[2, col - 1]).strip().lower()
                keyword = str(df.iloc[3, col - 1]).strip()
                payload = self.create_payload(device, language_code,
                                              target_url, keyword)
                print(payload)
                response = self.send_request(payload)
                result = self.parse_response(response)
                ws.cell(row=next_row, column=col, value=result)
            except Exception as e:
                ws.cell(row=next_row, column=col, value=f"Error: {e}")

        print("Saving in file...")
        try:
            wb.save(self.file_name)
            print(f"File saved: {self.file_name}")
        except Exception as e:
            print(f"Error while saving: {e}")


if __name__ == '__main__':
    url = "https://api.dataforseo.com/v3/serp/google/organic/live/advanced"
    filename = 'position_report.xlsx'
    auth = os.getenv("AUTHORIZATION")

    headers = {
        "Authorization": auth,
        "Content-Type": "application/json"
    }

    job_parse = PositionExtractor(file_name=filename, url_service=url,
                                  headers=headers)
    job_parse.get_site_position_in_file()
